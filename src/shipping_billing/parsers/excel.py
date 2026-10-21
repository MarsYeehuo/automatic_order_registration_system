from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Any

from ..models import ContentBlock, NormalizedDocument, SourceType
from ..utils import sha256_file
from .base import ParseContext, build_document, sniff_source_type

HEADER_TOKENS = {
    "charge",
    "charges",
    "description",
    "item",
    "fee",
    "fees",
    "amount",
    "currency",
    "cur",
    "qty",
    "quantity",
    "rate",
    "unitprice",
    "unit price",
    "payment",
    "term",
    "remark",
    "费用",
    "金额",
    "币种",
    "数量",
    "单价",
    "付款方式",
}


def cell_to_str(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return format(value, "f")
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return format(value, ".10g")
    return str(value).strip()


def _detect_header(rows: list[list[str]]) -> int | None:
    best_index = None
    best_score = 0
    for index, row in enumerate(rows[:30]):
        cells = {cell.strip().lower().replace("_", " ") for cell in row if cell.strip()}
        score = sum(1 for token in HEADER_TOKENS if any(token in cell for cell in cells))
        if score > best_score:
            best_score = score
            best_index = index
    return best_index if best_score >= 2 else None


def _trim_rows(rows: list[list[str]]) -> list[list[str]]:
    trimmed: list[list[str]] = []
    for row in rows:
        values = list(row)
        while values and values[-1] == "":
            values.pop()
        if values:
            trimmed.append(values)
    return trimmed


def _read_xlsx(path: Path) -> tuple[dict[str, list[list[str]]], list[str]]:
    from openpyxl import load_workbook

    notes: list[str] = []
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ValueError(f"xlsx 读取失败: {type(exc).__name__}: {exc}") from exc
    sheets: dict[str, list[list[str]]] = {}
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            values = [cell_to_str(value) for value in row[:100]]
            rows.append(values)
            if len(rows) >= 5000:
                notes.append(f"{sheet.title}: 超过 5000 行，已截断")
                break
        sheets[sheet.title] = _trim_rows(rows)
    workbook.close()
    return sheets, notes


def _read_xls(path: Path) -> tuple[dict[str, list[list[str]]], list[str]]:
    import xlrd

    notes: list[str] = []
    try:
        workbook = xlrd.open_workbook(str(path), on_demand=True)
    except Exception as exc:
        raise ValueError(f"xls 读取失败: {type(exc).__name__}: {exc}") from exc
    sheets: dict[str, list[list[str]]] = {}
    for sheet in workbook.sheets():
        rows: list[list[str]] = []
        for row_index in range(min(sheet.nrows, 5000)):
            values = [
                cell_to_str(sheet.cell_value(row_index, col_index))
                for col_index in range(min(sheet.ncols, 100))
            ]
            rows.append(values)
        if sheet.nrows > 5000:
            notes.append(f"{sheet.name}: 超过 5000 行，已截断")
        sheets[sheet.name] = _trim_rows(rows)
    workbook.release_resources()
    return sheets, notes


def parse_excel(path: Path, ctx: ParseContext, parent_document_id: str | None) -> NormalizedDocument:
    digest = sha256_file(path)
    source_type = sniff_source_type(path)
    blocks: list[ContentBlock] = []
    metadata: dict = {"parser": source_type.value}
    parse_error: str | None = None
    try:
        if source_type == SourceType.XLS:
            sheets, notes = _read_xls(path)
        else:
            sheets, notes = _read_xlsx(path)
        metadata["notes"] = notes
        sheet_meta: list[dict] = []
        for name, rows in sheets.items():
            if not rows:
                continue
            header_index = _detect_header(rows)
            block = ContentBlock(
                kind="table",
                location=f"sheet:{name},rows:1-{len(rows)}",
                rows=rows,
                metadata={
                    "source": "excel",
                    "sheet": name,
                    "header_row": header_index + 1 if header_index is not None else None,
                },
            )
            blocks.append(block)
            sheet_meta.append(
                {
                    "name": name,
                    "rows": len(rows),
                    "columns": max((len(row) for row in rows), default=0),
                    "header_row": header_index + 1 if header_index is not None else None,
                }
            )
        metadata["sheets"] = sheet_meta
        if not blocks:
            parse_error = "Excel 中没有可用工作表或数据"
    except Exception as exc:
        parse_error = str(exc)

    return build_document(
        path=path,
        source_type=source_type,
        sha256=digest,
        blocks=blocks,
        parent_document_id=parent_document_id,
        metadata=metadata,
        parse_error=parse_error,
    )
