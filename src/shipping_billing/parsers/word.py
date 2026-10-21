from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path

from ..models import ContentBlock, NormalizedDocument, SourceType
from ..utils import sha256_file
from .base import ParseContext, build_document, sniff_source_type


def _extract_doc_text(path: Path) -> tuple[str | None, str | None]:
    for command in (
        [shutil.which("antiword") or "", "-w", "0", str(path)],
        [shutil.which("catdoc") or "", str(path)],
    ):
        if not command[0]:
            continue
        try:
            completed = subprocess.run(
                command,
                capture_output=True,
                timeout=30,
                check=False,
            )
            if completed.returncode == 0 and completed.stdout.strip():
                return completed.stdout.decode("utf-8", errors="replace"), None
        except (OSError, subprocess.SubprocessError):
            continue
    data = path.read_bytes()
    pieces = re.findall(rb"[\x20-\x7e\r\n\t]{8,}", data)
    text = b"\n".join(pieces).decode("utf-8", errors="replace")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if text:
        return text, "antiword/catdoc 不可用，使用二进制可打印字符串回退解析"
    return None, "无法解析 .doc 文件"


def parse_word(path: Path, ctx: ParseContext, parent_document_id: str | None) -> NormalizedDocument:
    digest = sha256_file(path)
    source_type = sniff_source_type(path)
    blocks: list[ContentBlock] = []
    metadata: dict = {"parser": "python-docx"}
    parse_error: str | None = None
    if source_type == SourceType.DOC:
        text, warning = _extract_doc_text(path)
        metadata["parser"] = "antiword/catdoc/strings"
        if warning:
            metadata["warning"] = warning
        if text:
            blocks.append(
                ContentBlock(kind="text", location="doc:body", text=text, metadata={"source": "doc"})
            )
        else:
            parse_error = warning or "无法解析 .doc 文件"
    else:
        try:
            from docx import Document

            document = Document(str(path))
            paragraphs = [p.text.strip() for p in document.paragraphs if p.text and p.text.strip()]
            if paragraphs:
                blocks.append(
                    ContentBlock(
                        kind="text",
                        location="docx:paragraphs",
                        text="\n".join(paragraphs),
                        metadata={"source": "docx"},
                    )
                )
            table_meta: list[dict] = []
            for index, table in enumerate(document.tables, start=1):
                rows: list[list[str]] = []
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    while values and not values[-1]:
                        values.pop()
                    if values:
                        rows.append(values)
                if rows:
                    blocks.append(
                        ContentBlock(
                            kind="table",
                            location=f"docx:table:{index}",
                            rows=rows,
                            metadata={"source": "docx", "table_index": index},
                        )
                    )
                    table_meta.append({"index": index, "rows": len(rows)})
            metadata["tables"] = table_meta
            metadata["paragraphs"] = len(paragraphs)
            if not blocks:
                parse_error = "Word 文档中没有可用文本或表格"
        except Exception as exc:
            parse_error = f"docx 解析失败: {type(exc).__name__}: {exc}"

    return build_document(
        path=path,
        source_type=source_type,
        sha256=digest,
        blocks=blocks,
        parent_document_id=parent_document_id,
        metadata=metadata,
        parse_error=parse_error,
    )
